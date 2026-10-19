#!/usr/bin/env python3
"""
Daily updater for the Economy at a Glance page.

Pulls the latest figures from Reserve Bank of Australia statistical tables
(which republish ABS data) and rewrites the matching boxes in data/*.json.

Rules that keep the page safe:
  * A figure is only replaced when the RBA has a NEWER observation than the
    one already shown (meta.json -> auto[id].asOf). It never goes backwards,
    so a figure Claude's weekly check filled from a fresher source is kept.
  * Values outside a plausible range are skipped and reported.
  * If a table can't be downloaded, its boxes are left exactly as they are.

Usage:  python3 scripts/update_rba.py            (fetch from rba.gov.au)
        python3 scripts/update_rba.py --fixtures DIR   (read DIR/<table>-data.csv)
"""
import argparse
import csv
import io
import json
import sys
import urllib.request
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
UA = "Mozilla/5.0 (compatible; economy-at-a-glance updater; +https://github.com/phillippain/HSC-Economics)"
CSV_URL = "https://www.rba.gov.au/statistics/tables/csv/{t}-data.csv"
XLSX_URL = {  # fallback if the CSV is unavailable
    "a2": "https://www.rba.gov.au/statistics/tables/xls/a02hist.xlsx",
    "e2": "https://www.rba.gov.au/statistics/tables/xls/e02hist.xlsx",
    "h1": "https://www.rba.gov.au/statistics/tables/xls/h01hist.xlsx",
    "h2": "https://www.rba.gov.au/statistics/tables/xls/h02hist.xlsx",
    "h4": "https://www.rba.gov.au/statistics/tables/xls/h04hist.xlsx",
    "h5": "https://www.rba.gov.au/statistics/tables/xls/h05hist.xlsx",
    "i1": "https://www.rba.gov.au/statistics/tables/xls/i01hist.xlsx",
}
MINUS = "−"
warnings = []


def warn(msg):
    warnings.append(msg)
    print(f"::warning::{msg}" if not sys.stdout.isatty() else f"WARNING: {msg}")


# ---------------------------------------------------------------- reading RBA tables
def parse_date(s):
    if isinstance(s, datetime):
        return s.date()
    if isinstance(s, date):
        return s
    s = (s or "").strip()
    for fmt in ("%d-%b-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%y", "%b-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    v = str(v).strip().replace(",", "").replace(MINUS, "-")
    try:
        return float(v)
    except ValueError:
        return None


# Used only if a table ever arrives without its "Series ID" row.
TITLE_TO_ID = {
    "A$1=USD": "FXRUSD", "Trade-weighted Index May 1970 = 100": "FXRTWI",
    "Unemployment rate": "GLFSURSA", "Participation rate": "GLFSPRSA",
    "Year-ended wage growth": "GWPIYP", "Household saving ratio": "GGDPICHDISR",
    "Household debt to income": "BHFDDIT", "Year-ended real GDP growth": "GGDPCVGDPY",
    "Terms of trade": "GOPITT", "Current account balance as a per cent of output": "HCAGSCPGDP",
    "Change in Cash Rate Target": "ARBAMPCCCR", "New Cash Rate Target": "ARBAMPCNCRT",
}


def rows_to_series(rows):
    """rows: list of lists (first column = labels/dates). Returns {series_id: [(date, value), ...]}."""
    ids = None
    title_ids = None
    out = {}
    for r in rows:
        if not r:
            continue
        head = str(r[0]).strip() if r[0] is not None else ""
        if ids is None:
            if head in ("Series ID", "Mnemonic"):
                ids = [str(c).strip() if c is not None else "" for c in r]
            elif head == "Title":
                title_ids = [TITLE_TO_ID.get(str(c).strip(), "") if c is not None else "" for c in r]
            if ids is None and not (title_ids and parse_date(r[0])):
                continue
            ids = ids or title_ids
            for sid in ids[1:]:
                if sid:
                    out.setdefault(sid, [])
            if not parse_date(r[0]):
                continue
        d = parse_date(r[0])
        if not d:
            continue
        for i, sid in enumerate(ids[1:], start=1):
            if sid and i < len(r):
                v = parse_num(r[i])
                if v is not None:
                    out[sid].append((d, v))
    for sid in out:
        out[sid].sort()
    return out


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()


def load_table(t, fixtures=None):
    if fixtures:
        p = Path(fixtures) / f"{t}-data.csv"
        if not p.exists():
            warn(f"RBA table {t.upper()}: no fixture file")
            return {}
        return rows_to_series(list(csv.reader(io.StringIO(p.read_text(encoding="utf-8-sig")))))
    try:
        raw = fetch(CSV_URL.format(t=t))
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")
        s = rows_to_series(list(csv.reader(io.StringIO(text))))
        if s:
            return s
        raise ValueError("no 'Series ID' row in CSV")
    except Exception as e:  # noqa: BLE001
        if t not in XLSX_URL:
            warn(f"RBA table {t.upper()}: couldn't download ({e}); boxes left as they were")
            return {}
        try:
            import openpyxl  # only needed for the fallback

            wb = openpyxl.load_workbook(io.BytesIO(fetch(XLSX_URL[t])), read_only=True, data_only=True)
            ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
            return rows_to_series([list(r) for r in ws.iter_rows(values_only=True)])
        except Exception as e2:  # noqa: BLE001
            warn(f"RBA table {t.upper()}: couldn't download CSV ({e}) or Excel ({e2}); boxes left as they were")
            return {}


# ---------------------------------------------------------------- formatting
MON = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def sgn(x, dp=1, plus=True):
    s = f"{abs(x):.{dp}f}"
    if round(x, dp) < 0:
        return MINUS + s
    return ("+" + s) if plus and round(x, dp) > 0 else s


def pct(x, dp=1, plus=False):
    return sgn(x, dp, plus) + "%"


def month_label(d):
    return f"{MON[d.month - 1]} {d.year}"


def qtr_label(d):
    return f"{MON[d.month - 1]} qtr {d.year}"


def day_label(d):
    return f"{d.day} {MON[d.month - 1]} {d.year}"


def move(now, before, unit=" pts", dp=1):
    diff = round(now - before, dp)
    if diff == 0:
        return "steady"
    return ("up " if diff > 0 else "down ") + f"{abs(diff):.{dp}f}{unit}"


def last(series, n=1):
    return series[-n:] if series else []


def value_on_or_before(series, d):
    best = None
    for od, v in series:
        if od <= d:
            best = (od, v)
        else:
            break
    return best


# ---------------------------------------------------------------- the updates
PLAUSIBLE = {
    "GLFSURSA": (1, 20), "GLFSPRSA": (55, 75), "GWPIYP": (-5, 15), "GGDPICHDISR": (-10, 30),
    "BHFDDIT": (50, 300), "ARBAMPCNCRT": (0, 20), "ARBAMPCCCR": (-2, 2), "GOPITT": (10, 400),
    "FXRUSD": (0.3, 1.5), "FXRTWI": (30, 120), "HCAGSCPGDP": (-15, 15), "GGDPCVGDPY": (-15, 20),
}


def series(tables, t, sid):
    s = tables.get(t, {}).get(sid, [])
    if not s:
        if tables.get(t):
            warn(f"Series {sid} not found in RBA table {t.upper()}")
        return []
    lo, hi = PLAUSIBLE.get(sid, (-1e9, 1e9))
    bad = [v for _, v in s[-12:] if not lo <= v <= hi]
    if bad:
        warn(f"Series {sid} has implausible recent values {bad[:3]}; skipped")
        return []
    return s


def run(fixtures=None, today=None):
    today = today or date.today()
    glance = json.loads((DATA / "glance.json").read_text())
    score = json.loads((DATA / "scorecard.json").read_text())
    rates = json.loads((DATA / "policy_rates.json").read_text())
    meta_p = DATA / "meta.json"
    meta = json.loads(meta_p.read_text()) if meta_p.exists() else {}
    auto = meta.setdefault("auto", {})

    G = {r["id"]: r for r in glance}
    S = {r["id"]: r for r in score}
    changed = []

    def newer(key, d):
        prev = auto.get(key, {}).get("asOf")
        return prev is None or d.isoformat() > prev

    def set_glance(key, d, display, period, note, source=None):
        if not newer(key, d) or key not in G:
            return
        r = G[key]
        new = dict(r, display=display, period=period, note=note)
        if source:
            new["source"] = source
        if new != r:
            G[key].update(new)
            changed.append(f"{r['label']}: {display} ({period})")
        auto[key] = {"asOf": d.isoformat(), "by": "RBA statistical tables"}

    def set_score(key, d, fields):
        if not newer(key, d) or key not in S:
            return
        r = S[key]
        new = dict(r, **fields)
        if new != r:
            S[key].update(new)
            changed.append(f"Scorecard {r['measure']}: {fields.get('latest')}")
        auto[key] = {"asOf": d.isoformat(), "by": "RBA statistical tables"}

    tables = {t: load_table(t, fixtures) for t in ("h5", "h4", "h2", "h1", "e2", "a2", "f11.1", "i1")}

    # --- Labour force (RBA H5 = ABS 6202, seasonally adjusted)
    ur = series(tables, "h5", "GLFSURSA")
    if len(ur) >= 3:
        (d0, v0), (d1, v1), (d2, v2) = ur[-3], ur[-2], ur[-1]
        set_glance("unemp", d2, pct(v2), month_label(d2),
                   f"Seasonally adjusted; {move(v2, v1)} on {MON[d1.month-1]} (**{v1:.1f}%**).", "ABS 6202")
        set_score("unempsc", d2, {"latest": pct(v2), "latestPeriod": month_label(d2), "previous": pct(v1),
                                  "previousPeriod": month_label(d1), "earlier": pct(v0), "earlierPeriod": month_label(d0)})
    pr = series(tables, "h5", "GLFSPRSA")
    if len(pr) >= 3:
        (d0, v0), (d1, v1), (d2, v2) = pr[-3], pr[-2], pr[-1]
        mv = move(v2, v1)
        set_glance("partic", d2, pct(v2), month_label(d2),
                   (mv[0].upper() + mv[1:]) + f" on {MON[d1.month-1]} (**{v1:.1f}%**).", "ABS 6202")
        set_score("particsc", d2, {"latest": pct(v2), "latestPeriod": month_label(d2), "previous": pct(v1),
                                   "previousPeriod": month_label(d1), "earlier": pct(v0), "earlierPeriod": month_label(d0)})

    # --- Wages (RBA H4 = ABS 6345)
    wpi = series(tables, "h4", "GWPIYP")
    if len(wpi) >= 2:
        (d1, v1), (d2, v2) = wpi[-2], wpi[-1]
        set_glance("wpi", d2, pct(v2), qtr_label(d2),
                   f"Annual growth; **{v1:.1f}%** in the {qtr_label(d1)}. Compare with CPI inflation for real wages.", "ABS 6345")

    # --- Household saving ratio (RBA H2 = ABS 5206)
    sr = series(tables, "h2", "GGDPICHDISR")
    if len(sr) >= 2:
        (d1, v1), (d2, v2) = sr[-2], sr[-1]
        word = "Up" if v2 > v1 else ("Down" if v2 < v1 else "Unchanged")
        set_glance("saving", d2, pct(v2), qtr_label(d2), f"{word} from **{pct(v1)}** in the {qtr_label(d1)}.", "ABS 5206")

    # --- Household debt to income (RBA E2)
    dti = series(tables, "e2", "BHFDDIT")
    if len(dti) >= 2:
        (d1, v1), (d2, v2) = dti[-2], dti[-1]
        set_glance("debtinc", d2, f"{v2:.0f}%", qtr_label(d2),
                   f"Household debt as a share of annual disposable income; **{v1:.0f}%** in the {qtr_label(d1)}.", "RBA table E2")

    # --- GDP growth (RBA H1 = ABS 5206)
    gdp = series(tables, "h1", "GGDPCVGDPY")
    if len(gdp) >= 3:
        (d0, v0), (d1, v1), (d2, v2) = gdp[-3:]
        set_score("gdp", d2, {"latest": pct(v2), "latestPeriod": qtr_label(d2), "previous": pct(v1),
                              "previousPeriod": qtr_label(d1), "earlier": pct(v0), "earlierPeriod": qtr_label(d0)})

    # --- Terms of trade (RBA H1, index -> % change)
    tt = series(tables, "h1", "GOPITT")
    if len(tt) >= 6:
        q = [(tt[i][0], (tt[i][1] / tt[i - 1][1] - 1) * 100) for i in range(len(tt) - 3, len(tt))]
        yy = (tt[-1][1] / tt[-5][1] - 1) * 100
        (d0, c0), (d1, c1), (d2, c2) = q
        set_glance("tot", d2, pct(c2, plus=True), qtr_label(d2),
                   f"Quarter; **{pct(c1, plus=True)}** in the {qtr_label(d1)}; {pct(yy, plus=True)} over the year.", "ABS 5206")
        set_score("totsc", d2, {"latest": f"{pct(c2, plus=True)} ({pct(yy, plus=True)} y/y)", "latestPeriod": qtr_label(d2),
                                "previous": pct(c1, plus=True), "previousPeriod": qtr_label(d1),
                                "earlier": pct(c0, plus=True), "earlierPeriod": qtr_label(d0)})

    # --- Current account % of GDP (RBA I1 = ABS 5302)
    ca = series(tables, "i1", "HCAGSCPGDP")
    if len(ca) >= 2:
        (d1, v1), (d2, v2) = ca[-2], ca[-1]
        set_glance("capct", d2, pct(v2), qtr_label(d2),
                   f"**{pct(v1)}** in the {qtr_label(d1)}. Current account balance as a share of GDP (RBA table I1).", "ABS 5302")

    # --- Cash rate (RBA A2: one row per change)
    new_rate = series(tables, "a2", "ARBAMPCNCRT")
    chg = series(tables, "a2", "ARBAMPCCCR")
    if new_rate and chg:
        cmap = dict(chg)
        decisions = [(d, v, cmap.get(d)) for d, v in new_rate if cmap.get(d) not in (None, 0)]
        decisions = [x for x in decisions if x[0] <= today]
        if decisions:
            d2, rate, c = decisions[-1]
            this_year = [x for x in decisions if x[0].year == d2.year]
            net = sum(x[2] for x in this_year)
            set_glance("cash", d2, f"{rate:.2f}%", day_label(d2),
                       f"{'Up' if c > 0 else 'Down'} from **{rate - c:.2f}%**. Target band for inflation is 2–3%.", "RBA")
            n = len(this_year)
            set_glance("lastchg", d2, f"{sgn(c, 2)} pts", f"{'Raised' if c > 0 else 'Cut'} on {day_label(d2)}",
                       (f"Previous move **{sgn(decisions[-2][2], 2)} pts** on {decisions[-2][0].day} {MON[decisions[-2][0].month-1]}"
                        + (f" {decisions[-2][0].year}" if decisions[-2][0].year != d2.year else "") + ". "
                        if len(decisions) >= 2 else "") +
                       f"{n} change{'s' if n != 1 else ''} in {d2.year} "
                       f"({', '.join(MON[x[0].month-1] for x in this_year)}): {sgn(net, 2)} pts net.", "RBA media release")
            last3 = decisions[-3:][::-1] + [None] * 3
            def cell(x):
                return (f"{x[1]:.2f}% ({sgn(x[2], 2)})", day_label(x[0])) if x else ("—", "")
            (l, lp), (p, pp), (e, ep) = cell(last3[0]), cell(last3[1]), cell(last3[2])
            set_score("cashsc", d2, {"latest": l, "latestPeriod": lp, "previous": p, "previousPeriod": pp,
                                     "earlier": e, "earlierPeriod": ep})
            if newer("rate_rba", d2):
                for r in rates:
                    if r["bank"].startswith("Australia"):
                        r.update(rate=round(rate, 2), label=f"{rate:.2f}%", date=day_label(d2))
                auto["rate_rba"] = {"asOf": d2.isoformat(), "by": "RBA statistical tables"}

    # --- Exchange rates (RBA F11.1, daily 4pm AEST)
    usd = series(tables, "f11.1", "FXRUSD")
    if usd:
        d2, v2 = usd[-1]
        ago = value_on_or_before(usd, d2 - timedelta(days=365))
        prev_m = value_on_or_before(usd, d2.replace(day=1) - timedelta(days=1))
        if ago:
            ch = (v2 / ago[1] - 1) * 100
            set_glance("audchg", d2, pct(ch, plus=True), f"Year to {day_label(d2)}",
                       f"{v2*100:.2f} US cents on {day_label(d2)} vs **{ago[1]*100:.2f}** a year earlier.", "RBA")
            set_score("audsc", d2, {"latest": f"{v2*100:.2f}", "latestPeriod": day_label(d2),
                                    "previous": f"{prev_m[1]*100:.2f}" if prev_m else "—",
                                    "previousPeriod": day_label(prev_m[0]) if prev_m else "",
                                    "earlier": f"{ago[1]*100:.2f}", "earlierPeriod": day_label(ago[0]), "source": "RBA"})
    twi = series(tables, "f11.1", "FXRTWI")
    if twi:
        d2, v2 = twi[-1]
        ago = value_on_or_before(twi, d2 - timedelta(days=365))
        prev_m = value_on_or_before(twi, d2.replace(day=1) - timedelta(days=1))
        set_glance("twi", d2, f"{v2:.1f}", day_label(d2),
                   (f"**{prev_m[1]:.1f}** on {day_label(prev_m[0])}. " if prev_m else "") + "Trade-weighted index, May 1970 = 100.", "RBA")
        set_score("twisc", d2, {"latest": f"{v2:.1f}", "latestPeriod": day_label(d2),
                                "previous": f"{prev_m[1]:.1f}" if prev_m else "—",
                                "previousPeriod": day_label(prev_m[0]) if prev_m else "",
                                "earlier": f"{ago[1]:.1f}" if ago else "—",
                                "earlierPeriod": day_label(ago[0]) if ago else ""})

    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    # Only touch files when something changed, so quiet days make no commit.
    warnings_changed = meta.get("updaterWarnings", []) != warnings
    meta["updaterWarnings"] = warnings
    if changed:
        meta["updated"] = now
        (DATA / "glance.json").write_text(json.dumps(glance, ensure_ascii=False, indent=1) + "\n")
        (DATA / "scorecard.json").write_text(json.dumps(score, ensure_ascii=False, indent=1) + "\n")
        (DATA / "policy_rates.json").write_text(json.dumps(rates, ensure_ascii=False, indent=1) + "\n")
    if changed or warnings_changed or not meta_p.exists():
        meta_p.write_text(json.dumps(meta, ensure_ascii=False, indent=1) + "\n")

    print(f"{len(changed)} figure(s) updated" + (":" if changed else "."))
    for c in changed:
        print("  -", c)
    return changed


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--fixtures", help="read <table>-data.csv files from this folder instead of rba.gov.au")
    a = ap.parse_args()
    run(a.fixtures)
